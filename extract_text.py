import collections 
import collections.abc
import pptx
from PyPDF2 import PdfReader

print("--- PPTX ---")
prs = pptx.Presentation('GAI.pptx')
for slide in prs.slides:
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text)

print("\n--- PDF ---")
reader = PdfReader('Soft%20Engg%20and%20Pro%20Man_Sem%20I_P2023_Endsem_Dec2024_Data%20Science.pdf')
for i, page in enumerate(reader.pages):
    print(f"Page {i+1}:")
    print(page.extract_text())
